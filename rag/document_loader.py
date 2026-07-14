"""Asynchronous loaders for local knowledge-base documents."""

from __future__ import annotations

import asyncio
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable


@dataclass(frozen=True)
class LoadedDocument:
    """Raw document content and source metadata."""

    content: str
    metadata: dict[str, Any] = field(default_factory=dict)


class DocumentLoader:
    """Load TXT, Markdown, PDF, and DOCX documents without agent dependencies."""

    supported_extensions = {".txt", ".md", ".markdown", ".pdf", ".docx"}

    async def load(self, path: str | Path) -> list[LoadedDocument]:
        """Load one document, returning page-level records when available."""
        return await asyncio.to_thread(self._load_sync, Path(path))

    async def load_many(self, paths: Iterable[str | Path]) -> list[LoadedDocument]:
        """Load documents concurrently."""
        groups = await asyncio.gather(*(self.load(path) for path in paths))
        return [document for group in groups for document in group]

    def _load_sync(self, path: Path) -> list[LoadedDocument]:
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(path)
        extension = path.suffix.lower()
        if extension not in self.supported_extensions:
            raise ValueError(f"Unsupported document type: {extension}")

        base_metadata = {"source": str(path), "filename": path.name, "extension": extension}
        if extension in {".txt", ".md", ".markdown"}:
            return [LoadedDocument(path.read_text(encoding="utf-8"), base_metadata)]
        if extension == ".pdf":
            return self._load_pdf(path, base_metadata)
        return self._load_docx(path, base_metadata)

    @staticmethod
    def _load_pdf(path: Path, metadata: dict[str, Any]) -> list[LoadedDocument]:
        try:
            from pypdf import PdfReader
        except ImportError as error:  # pragma: no cover - installation guidance
            raise RuntimeError("PDF support requires pypdf.") from error
        reader = PdfReader(str(path))
        return [
            LoadedDocument(page.extract_text() or "", {**metadata, "page": index + 1})
            for index, page in enumerate(reader.pages)
            if page.extract_text()
        ]

    @staticmethod
    def _load_docx(path: Path, metadata: dict[str, Any]) -> list[LoadedDocument]:
        try:
            from docx import Document
        except ImportError as error:  # pragma: no cover - installation guidance
            raise RuntimeError("DOCX support requires python-docx.") from error
        content = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
        return [LoadedDocument(content, metadata)] if content.strip() else []
